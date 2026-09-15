from __future__ import annotations

import io
import json
from pathlib import Path
from zipfile import BadZipFile

import pandas as pd

from clasificacion_langchain.rag.schemas import KnowledgeDocument


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".docx"}


def _decode_content(content: bytes) -> str:
    return content.decode("utf-8", errors="ignore").strip()


def _read_csv_text(raw_text: str) -> str:
    if not raw_text.strip():
        return ""

    df = pd.read_csv(io.StringIO(raw_text))
    if df.empty:
        return ""

    if "text" in df.columns:
        series = df["text"].astype(str)
        return "\n".join(series.tolist()).strip()

    # Cada fila conserva el nombre de su columna ("precio_clp: 34990 | stock: 0"). Sin el
    # encabezado, un catalogo pierde el significado de cada valor y ni el retrieval ni el
    # LLM pueden responder "tienen stock?" o "cuanto cuesta?".
    columns = [str(column).strip() for column in df.columns]
    records: list[str] = []
    for row in df.fillna("").astype(str).itertuples(index=False):
        cells = [
            f"{column}: {value.strip()}"
            for column, value in zip(columns, row)
            if value and value.strip()
        ]
        if cells:
            records.append(" | ".join(cells))
    return "\n".join(records).strip()


def _read_json_text(raw_text: str) -> str:
    if not raw_text.strip():
        return ""

    payload = json.loads(raw_text)
    if isinstance(payload, list):
        blocks: list[str] = []
        for item in payload:
            if isinstance(item, dict) and "text" in item:
                blocks.append(str(item["text"]))
            else:
                blocks.append(json.dumps(item, ensure_ascii=False))
        return "\n".join(blocks).strip()

    if isinstance(payload, dict):
        if "text" in payload:
            return str(payload["text"]).strip()
        return json.dumps(payload, ensure_ascii=False).strip()

    return str(payload).strip()


def _read_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError(
            "Falta dependencia 'pypdf' para procesar archivos PDF"
        ) from exc

    reader = PdfReader(io.BytesIO(content))
    pages: list[str] = []
    for page in reader.pages:
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(page_text)
    return "\n\n".join(pages).strip()


def _read_docx_text(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "Falta dependencia 'python-docx' para procesar archivos DOCX"
        ) from exc

    try:
        document = Document(io.BytesIO(content))
    except BadZipFile as exc:
        raise ValueError("Archivo DOCX invalido") from exc

    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    return "\n".join(paragraphs).strip()


def load_text_content(filename: str | Path, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        raw_text = _decode_content(content)
        return raw_text
    if suffix == ".csv":
        raw_text = _decode_content(content)
        return _read_csv_text(raw_text)
    if suffix == ".json":
        raw_text = _decode_content(content)
        return _read_json_text(raw_text)
    if suffix == ".pdf":
        return _read_pdf_text(content)
    if suffix == ".docx":
        return _read_docx_text(content)
    return ""


def _load_file_content(path: Path) -> str:
    content = path.read_bytes()
    return load_text_content(path.name, content)


def load_company_documents(base_dir: str | Path) -> list[KnowledgeDocument]:
    root = Path(base_dir)
    if not root.exists():
        raise FileNotFoundError(
            f"No existe el directorio de conocimiento: {root}"
        )

    if not root.is_dir():
        raise ValueError(f"{root} debe ser un directorio")

    documents: list[KnowledgeDocument] = []

    for company_dir in sorted(root.iterdir()):
        if not company_dir.is_dir() or company_dir.name.startswith("."):
            continue

        company_id = company_dir.name
        for path in company_dir.rglob("*"):
            if not path.is_file():
                continue

            if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                continue

            text = _load_file_content(path)
            if not text:
                continue

            # Siempre con "/" para que el source sea estable entre Windows y Linux
            # (formato contractual: "<company_id>/<ruta relativa>").
            source = path.relative_to(root).as_posix()
            metadata = {
                "extension": path.suffix.lower(),
                "filename": path.name,
            }
            documents.append(
                KnowledgeDocument(
                    company_id=company_id,
                    source=source,
                    text=text,
                    metadata=metadata,
                )
            )

    if not documents:
        raise ValueError(
            "No se encontraron documentos validos. "
            "Usa estructura knowledge_base/<company_id>/*.txt|*.md|*.csv|*.json|*.pdf|*.docx"
        )

    return documents
